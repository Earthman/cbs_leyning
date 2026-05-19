import requests
import json
from datetime import datetime, timedelta
import sys
import argparse
from collections import defaultdict
from tenacity import retry, stop_after_attempt, wait_exponential
import pandas as pd
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

# ---------------------------------------------------------------------------
# Styling primitives
# ---------------------------------------------------------------------------

DEFAULT_FONT_NAME = 'Arial'
DEFAULT_FONT_SIZE = 11

# Background colors, expressed as the same 0-1 RGB the old gspread code used,
# kept here so the fallback layout reproduces the original sheets exactly.
GRAY = {'red': 0.9, 'green': 0.9, 'blue': 0.9}
ORANGE = {'red': 1.0, 'green': 0.8, 'blue': 0.6}
RED = {'red': 1.0, 'green': 0.8, 'blue': 0.8}
GREEN = {'red': 0.8, 'green': 1.0, 'blue': 0.8}
ALIYAH_COLORS = [
    {'red': 1.0, 'green': 1.0, 'blue': 0.8},
    {'red': 1.0, 'green': 0.8, 'blue': 1.0},
    {'red': 0.8, 'green': 1.0, 'blue': 1.0},
]


def gcolor_to_argb(gcolor):
    """Convert a {'red':0-1,'green':0-1,'blue':0-1} dict to an 'AARRGGBB' hex."""
    def chan(v):
        return format(max(0, min(255, round(v * 255))), '02X')
    return 'FF' + chan(gcolor.get('red', 0)) + chan(gcolor.get('green', 0)) + chan(gcolor.get('blue', 0))


def solid_fill(gcolor):
    argb = gcolor_to_argb(gcolor)
    return PatternFill(fill_type='solid', start_color=argb, end_color=argb)


def font(size=DEFAULT_FONT_SIZE, bold=False):
    return Font(name=DEFAULT_FONT_NAME, size=size, bold=bold)


def style_range(ws, cell_range, fill=None, cell_font=None, alignment=None):
    """Apply fill / font / alignment to every cell in an A1:F1 style range."""
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if fill is not None:
                cell.fill = fill
            if cell_font is not None:
                cell.font = cell_font
            if alignment is not None:
                cell.alignment = alignment


def write_cell(ws, row, col, value, size=DEFAULT_FONT_SIZE, bold=False,
               fill=None, center=False):
    """Write a value and apply the default Arial styling."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font(size=size, bold=bold)
    if fill is not None:
        cell.fill = fill
    if center:
        cell.alignment = Alignment(horizontal='center')
    return cell


def set_column_widths(ws, num_columns=6):
    """Match the pixel widths the template sheet used (px -> char units)."""
    # (column, pixel width) from the original Google Sheets template.
    pixel_widths = [
        ('A', 113), ('B', 233), ('C', 184),
        ('D', 184), ('E', 184), ('F', 442),
    ]
    for col, px in pixel_widths[:num_columns]:
        # openpyxl width is in character units; ~7px per unit, +5px padding.
        ws.column_dimensions[col].width = round((px - 5) / 7, 2)


def sanitize_sheet_name(name, used):
    """Make an Excel-legal, unique worksheet title (<=31 chars, no []:*?/\\)."""
    for ch in '[]:*?/\\':
        name = name.replace(ch, '-')
    name = name.strip() or 'Sheet'
    name = name[:31]
    candidate = name
    i = 2
    while candidate in used:
        suffix = f' ({i})'
        candidate = name[:31 - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


# ---------------------------------------------------------------------------
# Data helpers (unchanged from the Sheets implementation)
# ---------------------------------------------------------------------------

def find_marker_in_column(data, column_index, marker_text):
    """Search for marker text in a specific column. Returns 0-based row index."""
    for row_idx, row in enumerate(data):
        if column_index < len(row) and row[column_index] == marker_text:
            return row_idx
    return None


def int_to_roman(num):
    """Convert integer to Roman numeral."""
    roman_symbols = [
        ('M', 1000), ('CM', 900), ('D', 500), ('CD', 400),
        ('C', 100), ('XC', 90), ('L', 50), ('XL', 40),
        ('X', 10), ('IX', 9), ('V', 5), ('IV', 4), ('I', 1)
    ]
    result = ''
    for symbol, value in roman_symbols:
        while num >= value:
            result += symbol
            num -= value
    return result


def format_verse_range(aliyah):
    """Format verse range with verse count."""
    try:
        book = aliyah['k']

        start_parts = aliyah['b'].split(':')
        start_chapter = start_parts[0]
        start_verse = start_parts[1]

        end_parts = aliyah['e'].split(':')
        end_chapter = end_parts[0]
        end_verse = end_parts[1]

        if start_chapter == end_chapter:
            verse_range = f"{start_chapter}:{start_verse}-{end_verse}"
        else:
            verse_range = f"{start_chapter}:{start_verse}-{end_chapter}:{end_verse}"

        if 'v' in aliyah:
            return f"{book} {verse_range} ({aliyah['v']})"
        else:
            return f"{book} {verse_range}"

    except (KeyError, IndexError, AttributeError) as e:
        print(f"Error formatting verse range: {e}")
        print(f"Aliyah data: {aliyah}")
        return "Error formatting verse range"


def get_reading_type(name):
    """Determine the type of reading based on the name."""
    name_lower = name.lower()
    if 'fast' in name_lower or 'taanit' in name_lower:
        return 'fast_day'
    elif 'rosh chodesh' in name_lower:
        return 'rosh_chodesh'
    elif 'chol ha-moed' in name_lower or 'chol hamoed' in name_lower:
        return 'chol_hamoed'
    return 'regular'


def is_special_day(name):
    """Check if this is a special day that should be included in minyan readings."""
    name_lower = name.lower()
    return any(term in name_lower for term in [
        'rosh chodesh',
        'chol ha-moed',
        'chol hamoed',
        'fast',
        'taanit'
    ])


def load_page_numbers(csv_path):
    """Load page numbers from CSV file."""
    df = pd.read_csv(csv_path)
    if 'Haftara verses' in df.columns:
        df = df.rename(columns={'Haftara verses': 'Haftarah verses'})
    return df.set_index('Parsha').to_dict('index')


@retry(stop=stop_after_attempt(3), wait=wait_exponential(multiplier=1, min=4, max=10))
def get_leyning(start_date, end_date, verbose=False):
    """Fetch leyning data from HebCal API with retry logic."""
    url = f"https://www.hebcal.com/leyning?cfg=json&start={start_date}&end={end_date}"
    if verbose:
        print(f"Fetching data from {url}")
    response = requests.get(url)
    response.raise_for_status()
    return response.json()


# ---------------------------------------------------------------------------
# Book lookup (--book): find the date range of the next reading of a book
# ---------------------------------------------------------------------------

CANONICAL_BOOKS = ['Genesis', 'Exodus', 'Leviticus', 'Numbers', 'Deuteronomy']

BOOK_ALIASES = {
    'genesis': 'Genesis', 'bereshit': 'Genesis', 'bereishit': 'Genesis',
    "b'reshit": 'Genesis', 'breshit': 'Genesis',
    'exodus': 'Exodus', 'shemot': 'Exodus', 'shmot': 'Exodus',
    "sh'mot": 'Exodus',
    'leviticus': 'Leviticus', 'vayikra': 'Leviticus', 'vayikrah': 'Leviticus',
    'numbers': 'Numbers', 'bamidbar': 'Numbers', 'bemidbar': 'Numbers',
    "b'midbar": 'Numbers', 'bmidbar': 'Numbers',
    'deuteronomy': 'Deuteronomy', 'devarim': 'Deuteronomy',
    "d'varim": 'Deuteronomy', 'dvarim': 'Deuteronomy',
}


def canonical_book(name):
    """Map an English or transliterated book name to its canonical English
    name (as used in HebCal aliyah 'k' fields). Returns None if unknown."""
    if not name:
        return None
    return BOOK_ALIASES.get(name.strip().lower())


def _shabbat_parsha_book(item):
    """Return the Torah book of a regular weekly Shabbat parsha reading, or
    None for weekday previews, festivals, and other special days.

    A regular weekly parsha has a non-null `parshaNum`, no `weekday` flag,
    and a full kriyah whose first aliyah names the book."""
    if item.get('parshaNum') is None:
        return None
    if 'weekday' in item:
        return None
    fk = item.get('fullkriyah') or {}
    a1 = fk.get('1')
    if a1 is None and fk:
        a1 = next(iter(fk.values()))
    return a1.get('k') if a1 else None


def resolve_book_range(items, target_book, today_str):
    """Find the next reading of target_book on or after today_str.

    Returns (start_date, end_date, parshas, closed) where parshas is a list
    of (date, name) for the Shabbat readings in that book's run, and closed
    is True if a different book follows the run in the supplied items (so the
    run is known to be complete). Returns None if no upcoming reading is
    found in the supplied items.

    "Next reading" means the next time the book starts fresh: the first
    Shabbat parsha of target_book (on/after today) whose preceding weekly
    parsha belonged to a different book. The run then continues, skipping
    any interleaved festival weeks, until a different book begins.

    Items dated before today_str are still used for boundary detection (to
    tell a fresh start from a mid-book week) but never become the start.
    """
    weekly = []
    for it in items:
        book = _shabbat_parsha_book(it)
        if book:
            weekly.append((it['date'], book, it['name']['en']))
    weekly.sort(key=lambda x: x[0])

    start_idx = None
    for i, (date, book, _name) in enumerate(weekly):
        if (book == target_book and date >= today_str
                and (i == 0 or weekly[i - 1][1] != target_book)):
            start_idx = i
            break
    if start_idx is None:
        return None

    end_idx = start_idx
    while end_idx + 1 < len(weekly) and weekly[end_idx + 1][1] == target_book:
        end_idx += 1

    closed = (end_idx + 1 < len(weekly)
              and weekly[end_idx + 1][1] != target_book)
    parshas = [(d, n) for d, _b, n in weekly[start_idx:end_idx + 1]]
    return weekly[start_idx][0], weekly[end_idx][0], parshas, closed


def _next_day(date_str):
    return (datetime.strptime(date_str, '%Y-%m-%d')
            + timedelta(days=1)).strftime('%Y-%m-%d')


def find_next_book_reading(book, today_str, verbose=False, max_chunks=6):
    """Page through the HebCal API to find the next complete reading of book.

    HebCal's leyning endpoint caps each response at ~6 months, so a single
    wide request is not enough. Fetch ~175-day chunks (starting a few weeks
    before today, so book boundaries are visible) and stop as soon as the
    target book's run is bounded by the following book.

    Returns (items, start_date, end_date, parshas, closed) or None.
    """
    cursor = (datetime.strptime(today_str, '%Y-%m-%d')
              - timedelta(days=21)).strftime('%Y-%m-%d')
    collected = {}
    last_resolved = None

    for _ in range(max_chunks):
        chunk_end = (datetime.strptime(cursor, '%Y-%m-%d')
                     + timedelta(days=175)).strftime('%Y-%m-%d')
        d = get_leyning(cursor, chunk_end, verbose=verbose)
        items = d.get('items', [])
        if not items:
            break
        for it in items:
            collected[(it['date'], it['name']['en'], 'weekday' in it)] = it

        last_resolved = resolve_book_range(
            list(collected.values()), book, today_str)
        if last_resolved is not None and last_resolved[3]:  # closed
            break

        rng = d.get('range') or {}
        covered_end = rng.get('end') or max(it['date'] for it in items)
        nxt = _next_day(covered_end)
        if nxt <= cursor:  # no forward progress; give up
            break
        cursor = nxt

    if last_resolved is None:
        return None
    start_date, end_date, parshas, closed = last_resolved
    return list(collected.values()), start_date, end_date, parshas, closed


# ---------------------------------------------------------------------------
# Local .xlsx template loading
# ---------------------------------------------------------------------------

def _cell_style(cell):
    """Capture the fill / font-size / bold / alignment of a template cell."""
    style = {}
    fill = cell.fill
    if fill is not None and fill.fill_type == 'solid':
        rgb = getattr(fill.start_color, 'rgb', None)
        if isinstance(rgb, str) and len(rgb) == 8:
            style['fill_argb'] = rgb
    if cell.font is not None:
        if cell.font.size:
            style['size'] = cell.font.size
        style['bold'] = bool(cell.font.bold)
    if cell.alignment is not None and cell.alignment.horizontal:
        style['horizontal'] = cell.alignment.horizontal
    return style


def _read_template_sheet(ws, max_cols=6):
    """Return (values, styles) grids for a template worksheet."""
    values, styles = [], []
    for row in ws.iter_rows(min_col=1, max_col=max_cols):
        v_row, s_row = [], []
        for cell in row:
            v_row.append(cell.value if cell.value is not None else "")
            s_row.append(_cell_style(cell))
        values.append(v_row)
        styles.append(s_row)
    return values, styles


def load_template(template_path, verbose=False):
    """
    Load a local .xlsx template and detect dynamic dimensions.

    The template must have a 'Header' sheet and a 'Footer' sheet. Marker cells
    drive dynamic positioning so the layout can change without code edits:
      - "Torah(s) Scroll" in column A of Header marks the scroll row
      - "Reader"/"Aliyah"/"Hebrew Name(s)"/"Notes" in C-F marks the last
        header row (aliyot begin on the next row)

    Returns a dict with header/footer value + style grids and the detected
    indices, or raises if the file/markers are missing.
    """
    if verbose:
        print(f"Loading template: {template_path}", file=sys.stderr)

    wb = load_workbook(template_path, data_only=True)
    for required in ('Header', 'Footer'):
        if required not in wb.sheetnames:
            raise ValueError(f"Template is missing required sheet: '{required}'")

    header_values, header_styles = _read_template_sheet(wb['Header'])
    footer_values, footer_styles = _read_template_sheet(wb['Footer'])

    scroll_row = find_marker_in_column(header_values, 0, "Torah(s) Scroll")
    if scroll_row is None:
        raise ValueError("Template Header must contain 'Torah(s) Scroll' in column A")

    column_header_row = find_marker_in_column(header_values, 2, "Reader")
    if column_header_row is None:
        raise ValueError("Template Header must contain 'Reader' in column C")

    row = header_values[column_header_row]
    required_markers = {3: "Aliyah", 4: "Hebrew Name(s)", 5: "Notes"}
    for col_idx, marker in required_markers.items():
        if col_idx >= len(row) or row[col_idx] != marker:
            raise ValueError(
                f"Template Header row {column_header_row + 1} must contain "
                f"'{marker}' in column {get_column_letter(col_idx + 1)}")

    header_length = column_header_row + 1

    # Footer length = index of last non-empty row + 1 (keeps internal/leading
    # blank spacer rows, trims trailing blanks).
    footer_length = 0
    for idx, frow in enumerate(footer_values):
        if any(str(c).strip() for c in frow):
            footer_length = idx + 1

    if verbose:
        print(f"Template loaded: {template_path}", file=sys.stderr)
        print(f"  Header length: {header_length} rows", file=sys.stderr)
        print(f"  Torah Scroll row: {scroll_row + 1}", file=sys.stderr)
        print(f"  Column header row: {column_header_row + 1}", file=sys.stderr)
        print(f"  Footer length: {footer_length} rows", file=sys.stderr)

    return {
        'header_values': header_values,
        'header_styles': header_styles,
        'header_length': header_length,
        'scroll_row': scroll_row,
        'column_header_row': column_header_row,
        'footer_values': footer_values,
        'footer_styles': footer_styles,
        'footer_length': footer_length,
    }


def copy_block(ws, values, styles, start_row, num_rows, num_cols=6):
    """Copy a template value+style block onto the worksheet at start_row."""
    for i in range(num_rows):
        src_vals = values[i] if i < len(values) else []
        src_styles = styles[i] if i < len(styles) else []
        for c in range(num_cols):
            value = src_vals[c] if c < len(src_vals) else ""
            style = src_styles[c] if c < len(src_styles) else {}
            cell = ws.cell(row=start_row + i, column=c + 1, value=value)
            cell.font = font(size=style.get('size', DEFAULT_FONT_SIZE),
                             bold=style.get('bold', False))
            if 'fill_argb' in style:
                argb = style['fill_argb']
                cell.fill = PatternFill(fill_type='solid',
                                        start_color=argb, end_color=argb)
            if 'horizontal' in style:
                cell.alignment = Alignment(horizontal=style['horizontal'])


# ---------------------------------------------------------------------------
# Section writers
# ---------------------------------------------------------------------------

def write_header(ws, parsha_data, scroll_name="Gunther", template_data=None):
    """
    Write the header section. Returns the next available row (1-based).

    Uses the loaded template when available, otherwise falls back to the
    original hardcoded 14-row layout.
    """
    full_date = datetime.strptime(parsha_data['date'], '%Y-%m-%d')
    gregorian_date = full_date.strftime('%B %-d')
    previous_date = (full_date - timedelta(days=1)).strftime('%B %-d')

    hebrew_date_parts = parsha_data['hdate'].split()
    hebrew_date = f"{hebrew_date_parts[1]} {hebrew_date_parts[0]}"

    total_verses = 0
    parsha_verses = 0
    if 'fullkriyah' in parsha_data:
        for key, aliyah in parsha_data['fullkriyah'].items():
            verses = aliyah.get('v', 0)
            total_verses += verses
            if key != 'M':
                parsha_verses += verses

    special_shabbat = None
    if isinstance(parsha_data.get('reason'), dict):
        special_shabbat = parsha_data['reason'].get('haftara')
    if not special_shabbat and 'haft' in parsha_data:
        haft = parsha_data['haft']
        if isinstance(haft, dict):
            special_shabbat = haft.get('reason')
        elif isinstance(haft, list):
            for h in haft:
                if isinstance(h, dict) and 'reason' in h:
                    special_shabbat = h['reason']
                    break

    scroll_name_str = str(scroll_name) if scroll_name is not None else "Gunther"
    parsha_en = parsha_data['name']['en']

    def vt_header_link(scroll_cell):
        """Virtual Tikkun link whose scroll= value references the scroll-name
        cell, so editing that cell updates the link."""
        return (f'=hyperlink("https://myvirtualtikkun.com/?shul=cbssf'
                f'&view=both&scroll="&{scroll_cell}&"&parsha={parsha_en}", '
                f'"Virtual Tikkun")')

    musaf_formula = ('=if(ISNUMBER(SEARCH("Richman",$A$2)), '
                     '"RDR default", "RAR default")')
    verse_summary = f"Full kriyah - {total_verses} verses (parsha={parsha_verses})"

    if template_data:
        header_length = template_data['header_length']
        scroll_row = template_data['scroll_row']
        column_header_row = template_data['column_header_row']

        copy_block(ws, template_data['header_values'],
                   template_data['header_styles'], 1, header_length)

        # Overwrite dynamic cells (styling from the template is preserved
        # because copy_block already set the font/fill on these cells).
        ws['B1'] = parsha_en
        ws['D1'] = gregorian_date
        ws['E1'] = hebrew_date
        ws['B3'] = f"Kabbalat Shabbat {previous_date}"
        ws.cell(row=scroll_row + 1, column=2, value=scroll_name_str)
        ws.cell(row=scroll_row + 1, column=3,
                value=vt_header_link(f"$B${scroll_row + 1}"))
        ws.cell(row=column_header_row + 1, column=2, value=verse_summary)
        ws['C6'] = musaf_formula
        if special_shabbat:
            ws['D2'] = special_shabbat

        return header_length + 1

    # ---- Fallback: original hardcoded layout ----
    vt_link = vt_header_link("$B$13")
    header_data = [
        ["", parsha_en, "", gregorian_date, hebrew_date, ""],
        ["Rabbi Amanda Russell", "", "", special_shabbat or "", "", ""],
        ["Service leaders", f"Kabbalat Shabbat {previous_date}", "", "", "", ""],
        ["", "P'sukei D'zimrah", "", "", "", ""],
        ["", "Shacharit", "", "", "", ""],
        ["", "Musaf", musaf_formula, "", "", ""],
        ["", "Torah Service", "", "", "", ""],
        ["", "Gabbai", "Sam (default)", "", "", ""],
        ["", "Distribute honors", "Todd (default)", "", "", ""],
        ["", "Read announcements", "Jerilyn (default)", "", "", ""],
        ["Board hosts", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["Torah(s) Scroll", scroll_name_str, vt_link, "", "", ""],
        ["", verse_summary, "Reader", "Aliyah", "Hebrew Name(s)", "Notes"],
    ]
    for r, row_vals in enumerate(header_data, start=1):
        size = 24 if r == 1 else 14 if r == 2 else DEFAULT_FONT_SIZE
        for c, value in enumerate(row_vals, start=1):
            write_cell(ws, r, c, value, size=size)

    style_range(ws, 'A3:A3', fill=solid_fill(GRAY))
    style_range(ws, 'A11:A11', fill=solid_fill(GRAY))
    style_range(ws, 'A13:B13', fill=solid_fill(ORANGE))
    style_range(ws, 'A14:F14', fill=solid_fill(GRAY))

    return 15


def write_aliyot(ws, fullkriyah, parsha_data, start_row,
                 page_numbers=None, scroll_cell="$B$13"):
    """Write the aliyot section starting at start_row. Returns the next row.

    scroll_cell is the absolute reference to the cell holding the scroll
    name (B13 in the standard layout); the Virtual Tikkun links reference it
    so changing the scroll name updates every aliyah link automatically.
    """
    if not fullkriyah:
        return start_row

    parsha_name = parsha_data['name']['en'] if parsha_data else ""

    row = start_row
    color_index = 0

    def emit(label_formula, display_text, verse_info):
        nonlocal row, color_index
        fill = solid_fill(ALIYAH_COLORS[color_index])
        a = write_cell(ws, row, 1,
                       label_formula if label_formula else display_text,
                       fill=fill, center=True)
        if label_formula:
            a.alignment = Alignment(horizontal='center')
        write_cell(ws, row, 2, verse_info, fill=fill)
        write_cell(ws, row, 3, "", fill=fill)
        for c in range(4, 7):
            write_cell(ws, row, c, "")
        color_index = (color_index + 1) % 3
        row += 1

    for key in sorted(fullkriyah.keys()):
        if key == 'M':
            continue
        aliyah = fullkriyah[key]
        aliyah_num = int(key) if key.isdigit() else key
        display_num = (int_to_roman(int(aliyah_num))
                       if isinstance(aliyah_num, int) else aliyah_num)
        verse_info = format_verse_range(aliyah)
        vt_link = (f'=hyperlink("https://myvirtualtikkun.com/?shul=cbssf'
                   f'&scroll="&{scroll_cell}&"&parsha={parsha_name}'
                   f'&aliyah=A{aliyah_num}", "{display_num}")')
        emit(vt_link, display_num, verse_info)

    if 'M' in fullkriyah:
        maftir = fullkriyah['M']
        verse_info = format_verse_range(maftir)
        vt_link = (f'=hyperlink("https://myvirtualtikkun.com/?shul=cbssf'
                   f'&scroll="&{scroll_cell}&"&parsha={parsha_name}'
                   f'&aliyah=M", "Maf")')
        emit(vt_link, "Maf", verse_info)

    if parsha_data:
        verse_info = ""
        if page_numbers and pd.notna(page_numbers.get('Haftarah verses')):
            verse_info = page_numbers['Haftarah verses']
        elif 'haft' in parsha_data:
            haftarah_parts = parsha_data['haft']
            if isinstance(haftarah_parts, list):
                verse_parts = []
                total = 0
                for part in haftarah_parts:
                    verse_parts.append(f"{part['b']}-{part['e']}")
                    total += part['v']
                book = haftarah_parts[0]['k']
                verse_info = f"{book} {', '.join(verse_parts)} ({total})"
            else:
                part = haftarah_parts
                verse_info = f"{part['k']} {part['b']}-{part['e']} ({part['v']})"
        emit(None, "Haf", verse_info)

    return row


def write_footer(ws, start_row, page_numbers=None, template_data=None):
    """Write the footer section starting at start_row."""
    if page_numbers:
        torah_page = (f"Torah page {str(int(page_numbers['Torah Page']))}"
                      if pd.notna(page_numbers.get('Torah Page')) else "Torah page")
        haftarah_page = (f"Haftarah page {str(int(page_numbers['Haftarah Page']))}"
                         if pd.notna(page_numbers.get('Haftarah Page'))
                         else "Haftarah page")
    else:
        torah_page = "Torah page"
        haftarah_page = "Haftarah page"

    if template_data:
        footer_values = template_data['footer_values']
        footer_styles = template_data['footer_styles']
        footer_length = template_data['footer_length']

        copy_block(ws, footer_values, footer_styles, start_row, footer_length)

        # Page numbers go in column D, on the two rows after the "Etz Hayyim"
        # honors header (detected so template edits stay safe).
        etz_idx = None
        for idx in range(footer_length):
            row_vals = footer_values[idx] if idx < len(footer_values) else []
            if len(row_vals) > 3 and str(row_vals[3]).strip() == "Etz Hayyim":
                etz_idx = idx
                break
        if etz_idx is None:
            etz_idx = 1  # original layout: blank row, then honors header
        ws.cell(row=start_row + etz_idx + 1, column=4, value=torah_page)
        ws.cell(row=start_row + etz_idx + 2, column=4, value=haftarah_page)
        return

    # ---- Fallback: original hardcoded layout ----
    footer_data = [
        ["", "", "", "", "", ""],
        ["", "Honors", "", "Etz Hayyim", "", ""],
        ["P'ticha 1", "", "", torah_page, "", ""],
        ["P'ticha 2", "", "", haftarah_page, "", ""],
        ["Hagbah", "", "", "", "", ""],
        ["G'lilah", "", "", "", "", ""],
        ["Prayer for Country", "", "", "", "", ""],
        ["Prayer for Israel", "", "", "", "", ""],
        ["Prayer for Peace", "", "", "", "", ""],
        ["Anim Zmerot", "", "", "", "", ""],
        ["Adon Olam", "", "", "", "", ""],
    ]
    for i, row_vals in enumerate(footer_data):
        for c, value in enumerate(row_vals, start=1):
            write_cell(ws, start_row + i, c, value)

    gray = solid_fill(GRAY)
    header_row = start_row + 1
    for col in (1, 2, 4):
        ws.cell(row=header_row, column=col).fill = gray


def write_minyan(ws, parsha_data):
    """Write weekday Torah readings and special days to the Minyan sheet."""
    set_column_widths(ws, num_columns=4)

    readings = []
    for item in parsha_data['items']:
        if not ('weekday' in item
                or ('fullkriyah' in item and is_special_day(item['name']['en']))):
            continue
        readings.append({
            'readings': item.get('weekday', item.get('fullkriyah', {})),
            'parsha_name': item['name']['en'],
            'date': item['date'],
            'hdate': item['hdate'],
            'type': get_reading_type(item['name']['en']),
        })

    readings.sort(key=lambda x: x['date'])
    if not readings:
        return

    all_rows = []
    header_rows = []
    for reading_info in readings:
        date_obj = datetime.strptime(reading_info['date'], '%Y-%m-%d')
        secular_date = date_obj.strftime('%b %d')
        hebrew_date = ' '.join(reading_info['hdate'].split()[:-1])

        header_rows.append(len(all_rows))
        all_rows.append([secular_date, hebrew_date,
                         reading_info['parsha_name'], date_obj.strftime('%A')])

        for aliyah_num, reading in reading_info['readings'].items():
            if aliyah_num != 'M':
                roman_num = (int_to_roman(int(aliyah_num))
                             if aliyah_num.isdigit() else aliyah_num)
                all_rows.append([roman_num, format_verse_range(reading), '', ''])

        all_rows.append(['', '', '', ''])

    for r, row_vals in enumerate(all_rows, start=1):
        for c, value in enumerate(row_vals, start=1):
            write_cell(ws, r, c, value)

    for reading_info, header_row in zip(readings, header_rows):
        if reading_info['type'] == 'fast_day':
            bg = RED
        elif reading_info['type'] in ('rosh_chodesh', 'chol_hamoed'):
            bg = GREEN
        else:
            bg = GRAY

        hr = header_row + 1  # 1-based
        style_range(ws, f'A{hr}:D{hr}', fill=solid_fill(bg),
                    cell_font=font(bold=True),
                    alignment=Alignment(horizontal='center'))

        start = header_row + 2
        end = start
        while end < len(all_rows) and (end == start or any(all_rows[end - 1])):
            if all_rows[end - 1][0]:
                ws.cell(row=end, column=1).alignment = Alignment(horizontal='center')
            end += 1


# ---------------------------------------------------------------------------
# Workbook builder
# ---------------------------------------------------------------------------

def build_workbook(data, output_path, test_mode=False, page_numbers=None,
                   scroll_name="Gunther", verbose=False, template_data=None):
    """Build the leyning workbook locally and save it to output_path."""
    parsha_data = defaultdict(list)
    for item in data['items']:
        parsha_name = item['name']['en']
        if not is_special_day(parsha_name):
            parsha_data[parsha_name].append(item)

    if test_mode and parsha_data:
        first_parsha = next(iter(parsha_data))
        parsha_data = {first_parsha: parsha_data[first_parsha]}
        if verbose:
            print(f"Test mode: Processing only parsha {first_parsha}")

    wb = Workbook()
    used_names = set()

    minyan_ws = wb.active
    minyan_ws.title = sanitize_sheet_name("Minyan", used_names)
    write_minyan(minyan_ws, data)

    if verbose:
        print(f"Minyan tab complete. Processing {len(parsha_data)} parshas...",
              file=sys.stderr)

    for parsha_name, items in parsha_data.items():
        if verbose:
            print(f"Processing {parsha_name}")

        parsha_instance = next(
            (item for item in items if 'fullkriyah' in item), items[0])

        ws = wb.create_sheet(title=sanitize_sheet_name(parsha_name, used_names))
        set_column_widths(ws, num_columns=6)

        parsha_pages = page_numbers.get(parsha_name) if page_numbers else None

        next_row = write_header(ws, parsha_instance,
                                scroll_name=scroll_name,
                                template_data=template_data)
        scroll_row_1based = (template_data['scroll_row'] + 1
                             if template_data else 13)
        next_row = write_aliyot(ws, parsha_instance.get('fullkriyah', {}),
                                parsha_instance, start_row=next_row,
                                page_numbers=parsha_pages,
                                scroll_cell=f"$B${scroll_row_1based}")
        write_footer(ws, start_row=next_row + 1,
                     page_numbers=parsha_pages,
                     template_data=template_data)

    wb.save(output_path)

    print(f"\n{'=' * 60}")
    print(f"WORKBOOK CREATED: {output_path}")
    print(f"  Minyan + {len(parsha_data)} parsha sheet(s)")
    print(f"{'=' * 60}")
    return output_path


# ---------------------------------------------------------------------------
# Default template generator
# ---------------------------------------------------------------------------

def make_template(path):
    """Create a starter template.xlsx (Header + Footer sheets) users can edit."""
    wb = Workbook()
    header = wb.active
    header.title = 'Header'

    # Static layout; dynamic cells (B1, D1, E1, D2, B3, C6, B13, C13, B14)
    # are intentionally blank and get filled in per parsha.
    header_rows = [
        ["", "", "", "", "", ""],
        ["Rabbi Amanda Russell", "", "", "", "", ""],
        ["Service leaders", "", "", "", "", ""],
        ["", "P'sukei D'zimrah", "", "", "", ""],
        ["", "Shacharit", "", "", "", ""],
        ["", "Musaf", "", "", "", ""],
        ["", "Torah Service", "", "", "", ""],
        ["", "Gabbai", "Sam (default)", "", "", ""],
        ["", "Distribute honors", "Todd (default)", "", "", ""],
        ["", "Read announcements", "Jerilyn (default)", "", "", ""],
        ["Board hosts", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["Torah(s) Scroll", "", "", "", "", ""],
        ["", "", "Reader", "Aliyah", "Hebrew Name(s)", "Notes"],
    ]
    for r, row_vals in enumerate(header_rows, start=1):
        size = 24 if r == 1 else 14 if r == 2 else DEFAULT_FONT_SIZE
        for c, value in enumerate(row_vals, start=1):
            write_cell(header, r, c, value, size=size)
    set_column_widths(header, num_columns=6)
    style_range(header, 'A3:A3', fill=solid_fill(GRAY))
    style_range(header, 'A11:A11', fill=solid_fill(GRAY))
    style_range(header, 'A13:B13', fill=solid_fill(ORANGE))
    style_range(header, 'A14:F14', fill=solid_fill(GRAY))

    footer = wb.create_sheet(title='Footer')
    footer_rows = [
        ["", "", "", "", "", ""],
        ["", "Honors", "", "Etz Hayyim", "", ""],
        ["P'ticha 1", "", "", "", "", ""],
        ["P'ticha 2", "", "", "", "", ""],
        ["Hagbah", "", "", "", "", ""],
        ["G'lilah", "", "", "", "", ""],
        ["Prayer for Country", "", "", "", "", ""],
        ["Prayer for Israel", "", "", "", "", ""],
        ["Prayer for Peace", "", "", "", "", ""],
        ["Anim Zmerot", "", "", "", "", ""],
        ["Adon Olam", "", "", "", "", ""],
    ]
    for r, row_vals in enumerate(footer_rows, start=1):
        for c, value in enumerate(row_vals, start=1):
            write_cell(footer, r, c, value)
    set_column_widths(footer, num_columns=6)
    gray = solid_fill(GRAY)
    for col in (1, 2, 4):
        footer.cell(row=2, column=col).fill = gray

    wb.save(path)
    print(f"Template written to {path}")
    return path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

DEFAULT_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                'template.xlsx')


def main():
    parser = argparse.ArgumentParser(
        description='Generate Torah-reading leyning sheets locally as .xlsx')
    parser.add_argument('start_date', nargs='?',
                        help='Start date in YYYY-MM-DD format')
    parser.add_argument('end_date', nargs='?',
                        help='End date in YYYY-MM-DD format')
    parser.add_argument('-b', '--book',
                        help='Generate the next reading of a Torah book '
                             '(e.g. Leviticus or Vayikra). Looks up the date '
                             'range automatically; no dates needed.')
    parser.add_argument('-v', '--verbose', action='store_true',
                        help='Enable verbose output')
    parser.add_argument('-s', '--sheet',
                        help='Output .xlsx path (e.g. leyning_5786.xlsx)')
    parser.add_argument('-t', '--test', action='store_true',
                        help='Test mode - only process first parsha')
    parser.add_argument('--pages', help='CSV file with page numbers')
    parser.add_argument('--scroll', help='Name of scroll (default is Gunther)')
    parser.add_argument('--template', default=DEFAULT_TEMPLATE,
                        help='Path to a local .xlsx template '
                             '(default: template.xlsx beside this script). '
                             'Falls back to the built-in layout if missing.')
    parser.add_argument('--json',
                        help='Read HebCal leyning JSON from a local file '
                             'instead of calling the API')
    parser.add_argument('--make-template', metavar='PATH',
                        help='Write a starter template .xlsx to PATH and exit')

    args = parser.parse_args()

    if args.make_template:
        make_template(args.make_template)
        return

    book = None
    if args.book:
        book = canonical_book(args.book)
        if not book:
            parser.error(
                f"unknown book '{args.book}'. Use one of: "
                + ", ".join(CANONICAL_BOOKS)
                + " (Hebrew names like Vayikra/Bamidbar also accepted).")
    elif not args.start_date or not args.end_date:
        parser.error("provide START_DATE and END_DATE, or --book BOOK "
                     "(or --make-template)")

    if book:
        today_str = datetime.now().strftime('%Y-%m-%d')
        if args.json:
            with open(args.json, 'r', encoding='utf-8') as f:
                src = json.load(f)
            resolved = resolve_book_range(src['items'], book, today_str)
            if resolved is None:
                print(f"Error: could not find an upcoming reading of {book} "
                      f"in {args.json}.", file=sys.stderr)
                sys.exit(1)
            start_date, end_date, parshas, closed = resolved
            items_pool = src['items']
        else:
            if args.verbose:
                print(f"Looking up next {book} reading from {today_str}...",
                      file=sys.stderr)
            found = find_next_book_reading(book, today_str,
                                           verbose=args.verbose)
            if found is None:
                print(f"Error: could not find an upcoming reading of {book}.",
                      file=sys.stderr)
                sys.exit(1)
            items_pool, start_date, end_date, parshas, closed = found

        if not closed:
            print(f"Warning: the reading range for {book} may be incomplete "
                  f"(no following book found in the available data).",
                  file=sys.stderr)

        # Restrict to this book's window so the workbook (including the
        # weekday Minyan readings) covers exactly this book.
        data = {'items': [it for it in items_pool
                          if start_date <= it['date'] <= end_date]}

        print(f"{book}: {start_date} ({parshas[0][1]}) -> "
              f"{end_date} ({parshas[-1][1]}), {len(parshas)} parshas")
        if args.verbose:
            for d, n in parshas:
                print(f"  {d}  {n}", file=sys.stderr)
    else:
        try:
            datetime.strptime(args.start_date, '%Y-%m-%d')
            datetime.strptime(args.end_date, '%Y-%m-%d')
        except ValueError:
            print("Error: Dates must be in YYYY-MM-DD format", file=sys.stderr)
            sys.exit(1)

        if args.json:
            with open(args.json, 'r', encoding='utf-8') as f:
                data = json.load(f)
        else:
            data = get_leyning(args.start_date, args.end_date,
                               verbose=args.verbose)

    page_numbers = None
    if args.pages:
        page_numbers = load_page_numbers(args.pages)

    scroll_name = args.scroll if args.scroll else "Gunther"

    template_data = None
    if args.template and os.path.exists(args.template):
        try:
            template_data = load_template(args.template, verbose=args.verbose)
            if args.verbose:
                print(f"Template mode enabled: {args.template}")
        except Exception as e:
            print(f"Warning: Could not load template '{args.template}': {e}",
                  file=sys.stderr)
            print("Falling back to the built-in hardcoded layout.",
                  file=sys.stderr)
            template_data = None
    elif args.verbose:
        print(f"No template at '{args.template}'; using built-in layout.",
              file=sys.stderr)

    output_path = args.sheet
    if not output_path and book:
        output_path = f"{book}.xlsx"

    if output_path:
        if not output_path.lower().endswith('.xlsx'):
            output_path += '.xlsx'
        build_workbook(data, output_path,
                       test_mode=args.test,
                       page_numbers=page_numbers,
                       scroll_name=scroll_name,
                       verbose=args.verbose,
                       template_data=template_data)
        print(f"\nData written to: {output_path}")
    else:
        print(json.dumps(data, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
